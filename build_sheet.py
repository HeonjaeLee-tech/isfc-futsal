# -*- coding: utf-8 -*-
"""ISFC 20주년 OBYB 풋살대회 운영 시트(.xlsx) 생성기.
구글 드라이브에 올려 'Google 스프레드시트로 저장'해서 쓰는 템플릿.
사이트(assets/js/app.js)는 각 시트 1행의 칸 이름으로 데이터를 읽는다.
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUT = sys.argv[1] if len(sys.argv) > 1 else "ISFC_풋살대회_운영시트.xlsx"

FONT = "Arial"
BURGUNDY = "741243"
MINT = "BEE0A9"
INPUT = "FFF6D5"   # 입력 칸
FIXED = "ECECEC"   # 고정 칸
EXAMPLE = "F4F4F4"
ROWS = 200         # 서식/드롭다운을 미리 깔아둘 행 수

thin = Side(style="thin", color="D6D6D6")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()
wb.remove(wb.active)


def header(ws, cols, notes=None, widths=None):
    for i, name in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", fgColor=BURGUNDY)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER
        if notes and name in notes:
            cm = Comment(notes[name], "ISFC 운영진")
            cm.width, cm.height = 260, 110
            c.comment = cm
    ws.row_dimensions[1].height = 26
    ws.freeze_panes = "A2"
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def body(ws, ncols, fixed_cols=(), rows=ROWS):
    """모든 데이터 칸을 '일반 텍스트'(@)로. 시간·숫자가 자동 변환되지 않게."""
    for r in range(2, rows + 2):
        for col in range(1, ncols + 1):
            c = ws.cell(row=r, column=col)
            c.number_format = "@"
            c.font = Font(name=FONT, size=11)
            c.alignment = Alignment(vertical="center", wrap_text=False)
            c.border = BORDER
            c.fill = PatternFill("solid", fgColor=FIXED if col in fixed_cols else INPUT)
        ws.row_dimensions[r].height = 22


def put(ws, r, values, fixed_cols=()):
    for col, v in enumerate(values, start=1):
        c = ws.cell(row=r, column=col, value=None if v is None else str(v))
        c.number_format = "@"
        if col in fixed_cols:
            c.font = Font(name=FONT, size=11, color="555555")


def example(ws, r, values):
    put(ws, r, values)
    for col in range(1, len(values) + 1):
        c = ws.cell(row=r, column=col)
        c.font = Font(name=FONT, size=11, italic=True, color="8A8A8A")
        c.fill = PatternFill("solid", fgColor=EXAMPLE)


def dropdown(ws, col_letter, options=None, formula=None, rows=ROWS, prompt=None, ref=None, start=2):
    f = formula if formula else '"' + ",".join(options) + '"'
    dv = DataValidation(type="list", formula1=f, allow_blank=True, showDropDown=False)
    dv.error = "목록에서 골라주세요"
    dv.errorTitle = "입력 확인"
    dv.showErrorMessage = True
    if prompt:
        dv.prompt = prompt
        dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(ref or f"{col_letter}{start}:{col_letter}{rows + 1}")


# ───────────────────────── 가이드 ─────────────────────────
g = wb.create_sheet("가이드")
g.sheet_properties.tabColor = BURGUNDY
g.column_dimensions["A"].width = 4
g.column_dimensions["B"].width = 22
g.column_dimensions["C"].width = 92
guide = [
    ("title", "ISFC 20주년 OBYB 풋살대회 운영 시트"),
    ("sub", "이 시트를 고치면 20~30초 안에 사이트에 반영돼요. 순위표, 대진표, 득점왕은 사이트가 알아서 계산합니다."),
    ("gap", ""),
    ("h", "처음 한 번만"),
    ("row", ("1", "이 파일을 구글 드라이브에 올린 뒤 [파일 → Google 스프레드시트로 저장]으로 변환")),
    ("row", ("2", "오른쪽 위 [공유] → 일반 액세스를 '링크가 있는 모든 사용자 · 뷰어'로 변경")),
    ("row", ("3", "주소창 링크에서 /d/ 와 /edit 사이 글자(시트 ID)를 사이트 config.js의 SHEET_ID에 붙여넣기")),
    ("gap", ""),
    ("h", "색깔 규칙"),
    ("row", ("노란 칸", "여기에 입력하세요")),
    ("row", ("회색 칸", "고정값이에요. 건드리지 마세요 (경기 번호, 단계, 팀코드 등)")),
    ("row", ("연회색 기울임", "(예시) 로 시작하는 줄은 사이트에 안 나와요. 형식만 참고하고 그대로 두거나 지워도 돼요")),
    ("row", ("주의", "1행 제목줄과 시트(탭) 이름은 절대 바꾸지 마세요. 사이트가 이 이름으로 읽어요")),
    ("gap", ""),
    ("h", "당일 구장 담당 (이것만 하면 돼요)"),
    ("row", ("경기 시작", "[경기] 시트에서 해당 번호의 상태를 '진행중'으로. 늦게 시작했으면 '시작' 칸도 실제 시각으로 고쳐야 사이트 경기 타이머가 맞아요")),
    ("row", ("골", "[경기] 홈골 / 원정골 숫자 수정 → [득점] 시트 맨 아래에 경기번호 + 선수 이름 한 줄")),
    ("row", ("경기 끝", "상태를 '종료'로. 그 순간 순위표와 토너먼트 대진이 자동으로 채워져요")),
    ("row", ("토너먼트 무승부", "승부차기 칸에 '홈-원정' 순서로 입력 (예: 4-3)")),
    ("row", ("자책골", "[득점] 메모 칸에 '자책골'이라고 쓰면 득점왕 집계에서 빠져요")),
    ("gap", ""),
    ("h", "공지 띄우기"),
    ("row", ("새 공지", "[공지] 시트 맨 아래 줄에 추가하면 사이트 맨 위 실시간 배너에 바로 떠요")),
    ("row", ("긴급", "중요도를 '긴급'으로 하면 배너가 버건디색으로 강조돼요")),
    ("row", ("내리기", "노출을 X로 바꾸면 사라져요 (지우지 않아도 됨)")),
    ("gap", ""),
    ("h", "시트별 역할"),
    ("row", ("공지", "실시간 공지 배너 + 공지 목록")),
    ("row", ("경기", "경기 일정, 점수, 진행 상태. 토너먼트의 'A조 2위', '21경기 승자'는 자동으로 실제 팀으로 바뀜")),
    ("row", ("득점", "골 넣은 사람 기록 → 득점왕 순위")),
    ("row", ("팀", "당일 제비뽑기 결과: 팀 이름, 조끼 색, 주장, 한마디 (팀코드 A1~B5는 고정)")),
    ("row", ("선수", "선수 명단. 팀코드를 적으면 팀 페이지에 자동 배치")),
    ("row", ("기념품", "텀블러, 키링 남은 수량. 배부수량만 틈틈이 올려주세요")),
    ("row", ("안주", "뒤풀이 안주 품목과 준비 상태")),
    ("row", ("뒤풀이", "1차, 2차 장소와 시간, 회비, 이동 방법")),
    ("row", ("설정", "팀 공개 여부(O/X)와 경기 규칙 문구")),
]
r = 2
for kind, val in guide:
    if kind == "title":
        c = g.cell(row=r, column=2, value=val)
        c.font = Font(name=FONT, bold=True, size=18, color=BURGUNDY)
        g.row_dimensions[r].height = 34
    elif kind == "sub":
        c = g.cell(row=r, column=2, value=val)
        c.font = Font(name=FONT, size=11, color="444444")
    elif kind == "h":
        for col in (2, 3):
            g.cell(row=r, column=col).fill = PatternFill("solid", fgColor=MINT)
        c = g.cell(row=r, column=2, value=val)
        c.font = Font(name=FONT, bold=True, size=12, color="1F3D12")
        g.row_dimensions[r].height = 24
    elif kind == "row":
        k, v = val
        a = g.cell(row=r, column=2, value=k)
        b = g.cell(row=r, column=3, value=v)
        a.font = Font(name=FONT, bold=True, size=11)
        b.font = Font(name=FONT, size=11)
        if k == "노란 칸":
            a.fill = PatternFill("solid", fgColor=INPUT)
        if k == "회색 칸":
            a.fill = PatternFill("solid", fgColor=FIXED)
        if k == "연회색 기울임":
            a.fill = PatternFill("solid", fgColor=EXAMPLE)
            a.font = Font(name=FONT, bold=True, italic=True, size=11, color="8A8A8A")
        if k == "주의":
            a.font = Font(name=FONT, bold=True, size=11, color=BURGUNDY)
            b.font = Font(name=FONT, bold=True, size=11, color=BURGUNDY)
        g.row_dimensions[r].height = 20
    r += 1
g.sheet_view.showGridLines = False

# ───────────────────────── 공지 ─────────────────────────
ws = wb.create_sheet("공지")
ws.sheet_properties.tabColor = "C0392B"
cols = ["시간", "내용", "중요도", "노출"]
header(ws, cols, notes={
    "시간": "자유 형식. 예: 10/31 14:05 또는 14:05",
    "내용": "한 줄로 짧게. 맨 아래 줄이 가장 최신 공지로 배너에 떠요",
    "중요도": "일반 / 중요 / 긴급. 긴급은 배너가 강조돼요",
    "노출": "O = 보이기, X = 숨기기",
}, widths=[16, 70, 10, 8])
body(ws, 4)
dropdown(ws, "C", ["일반", "중요", "긴급"])
dropdown(ws, "D", ["O", "X"])
put(ws, 2, ["10/25", "ISFC 20주년 OBYB 풋살대회 안내 페이지 오픈! 오시기 전에 주차 안내 먼저 확인해 주세요", "일반", "O"])

# ───────────────────────── 경기 ─────────────────────────
ws = wb.create_sheet("경기")
ws.sheet_properties.tabColor = BURGUNDY
cols = ["번호", "단계", "시작", "구장", "홈", "원정", "홈골", "원정골", "승부차기", "상태", "메모"]
header(ws, cols, notes={
    "번호": "고정. 득점 시트에서 이 번호로 경기를 가리켜요",
    "단계": "고정. A조 / B조 / 4강진출전 / 4강 / 3·4위전 / 결승",
    "시작": "시:분 (예: 14:15). 일정이 밀리면 여기만 고치면 돼요. 사이트 경기 타이머도 이 시각부터 흘러요",
    "홈": "조별리그는 팀코드(A1~B5). 토너먼트는 'A조 2위', '21경기 승자'처럼 두면 사이트가 실제 팀으로 바꿔줘요",
    "원정": "홈과 같은 규칙",
    "홈골": "숫자만. 경기 시작 전엔 비워두세요",
    "원정골": "숫자만",
    "승부차기": "토너먼트 무승부일 때만. 홈-원정 순서 (예: 4-3)",
    "상태": "예정 → 진행중 → 종료",
    "메모": "사이트에 작게 표시돼요 (선택)",
}, widths=[7, 12, 8, 8, 14, 14, 7, 7, 10, 9, 26])
FIXED_MATCH = (1, 2)
body(ws, len(cols), fixed_cols=FIXED_MATCH, rows=40)
dropdown(ws, "D", ["1구장", "2구장"], rows=40)
dropdown(ws, "J", ["예정", "진행중", "종료"], rows=40)

slots = ["13:15", "13:30", "13:45", "14:00", "14:15", "14:30", "14:45", "15:00", "15:15", "15:30"]
pairs = [(1, 2), (3, 4), (5, 1), (2, 3), (4, 5), (1, 3), (2, 4), (3, 5), (1, 4), (2, 5)]
rows = []
n = 1
for t, (a, b) in zip(slots, pairs):
    rows.append([n, "A조", t, "1구장", f"A{a}", f"A{b}", "", "", "", "예정", ""]); n += 1
    rows.append([n, "B조", t, "2구장", f"B{a}", f"B{b}", "", "", "", "예정", ""]); n += 1
rows += [
    [21, "4강진출전", "15:45", "1구장", "A조 2위", "B조 3위", "", "", "", "예정", ""],
    [22, "4강진출전", "15:45", "2구장", "B조 2위", "A조 3위", "", "", "", "예정", ""],
    [23, "4강", "16:00", "1구장", "A조 1위", "21경기 승자", "", "", "", "예정", ""],
    [24, "4강", "16:00", "2구장", "B조 1위", "22경기 승자", "", "", "", "예정", ""],
    [25, "3·4위전", "16:20", "1구장", "23경기 패자", "24경기 패자", "", "", "", "예정", ""],
    [26, "결승", "16:35", "1구장", "23경기 승자", "24경기 승자", "", "", "", "예정", "15분 단판"],
]
for i, row in enumerate(rows, start=2):
    put(ws, i, row, fixed_cols=FIXED_MATCH)
for col in (1, 3, 4, 7, 8, 9, 10):
    for i in range(2, 42):
        ws.cell(row=i, column=col).alignment = Alignment(horizontal="center", vertical="center")

# ───────────────────────── 득점 ─────────────────────────
ws = wb.create_sheet("득점")
ws.sheet_properties.tabColor = "E5B523"
cols = ["경기번호", "선수", "메모"]
header(ws, cols, notes={
    "경기번호": "[경기] 시트의 번호",
    "선수": "[선수] 시트에 적힌 이름 그대로 (목록에서 선택 가능)",
    "메모": "자책골이면 '자책골'이라고 쓰면 득점왕 집계에서 빠져요",
}, widths=[10, 18, 30])
body(ws, 3)
dropdown(ws, "A", [str(i) for i in range(1, 27)], start=3)
dropdown(ws, "B", formula="=선수!$A$3:$A$201", start=3)
example(ws, 2, ["(예시) 1", "(예시) 홍길동", "왼발 중거리"])

# ───────────────────────── 팀 ─────────────────────────
ws = wb.create_sheet("팀")
ws.sheet_properties.tabColor = "4E7A3A"
cols = ["팀코드", "팀명", "조끼색", "주장", "한마디"]
header(ws, cols, notes={
    "팀코드": "고정 (A1~A5는 A조, B1~B5는 B조)",
    "팀명": "비우면 사이트에 팀코드로 나와요",
    "조끼색": "목록에서 선택. 사이트 팀 색으로 쓰여요",
    "주장": "이름",
    "한마디": "팀 근황이나 각오 한 줄",
}, widths=[9, 18, 10, 12, 60])
body(ws, len(cols), fixed_cols=(1,), rows=10)
dropdown(ws, "C", ["빨강", "주황", "노랑", "연두", "초록", "하늘", "파랑", "남색", "보라", "분홍", "검정", "흰색", "회색"], rows=10)
for i, code in enumerate(["A1", "A2", "A3", "A4", "A5", "B1", "B2", "B3", "B4", "B5"], start=2):
    put(ws, i, [code, "", "", "", ""], fixed_cols=(1,))
    ws.cell(row=i, column=1).alignment = Alignment(horizontal="center", vertical="center")

# ───────────────────────── 선수 ─────────────────────────
ws = wb.create_sheet("선수")
ws.sheet_properties.tabColor = "4E7A3A"
cols = ["이름", "학번", "구분", "팀코드", "메모"]
header(ws, cols, notes={
    "이름": "득점 기록에도 이 이름 그대로 써요. 동명이인은 '김민수(05)'처럼 구분",
    "학번": "두 자리 (예: 05, 22)",
    "구분": "OB / YB",
    "팀코드": "A1~B5. 팀 매칭 후 입력",
    "메모": "포지션, 별명 등 (선택)",
}, widths=[16, 8, 8, 9, 30])
body(ws, len(cols))
dropdown(ws, "C", ["OB", "YB"], start=3)
dropdown(ws, "D", ["A1", "A2", "A3", "A4", "A5", "B1", "B2", "B3", "B4", "B5"], start=3)
example(ws, 2, ["(예시) 홍길동", "05", "OB", "A1", "주장, 왼발"])

# ───────────────────────── 기념품 ─────────────────────────
ws = wb.create_sheet("기념품")
ws.sheet_properties.tabColor = "A35E88"
cols = ["품목", "총수량", "배부수량", "수령장소", "메모"]
header(ws, cols, notes={
    "배부수량": "나눠준 개수. 틈틈이 올리면 사이트에 남은 수량이 떠요",
    "수령장소": "예: 풋살장 입구 접수대",
}, widths=[18, 9, 10, 24, 40])
body(ws, len(cols), rows=10)
put(ws, 2, ["20주년 텀블러", "200", "0", "", "풋살 또는 뒤풀이 참여 시 1인 1개"])
put(ws, 3, ["하냥이 키링", "200", "0", "", "풋살 또는 뒤풀이 참여 시 1인 1개"])

# ───────────────────────── 안주 ─────────────────────────
ws = wb.create_sheet("안주")
ws.sheet_properties.tabColor = "A35E88"
cols = ["메뉴", "수량", "상태", "메모"]
header(ws, cols, notes={
    "메뉴": "뒤풀이에 준비하는 안주 품목",
    "상태": "준비중 / 도착 / 세팅완료 / 소진",
}, widths=[20, 10, 10, 36])
body(ws, len(cols), rows=60)
dropdown(ws, "C", ["준비중", "도착", "세팅완료", "소진"], rows=60, start=3)
example(ws, 2, ["(예시) 치킨", "10마리", "도착", "반반"])

# ───────────────────────── 뒤풀이 ─────────────────────────
ws = wb.create_sheet("뒤풀이")
ws.sheet_properties.tabColor = "A35E88"
cols = ["항목", "내용"]
header(ws, cols, notes={"항목": "고정", "내용": "비워두면 사이트에 '공지 예정'으로 떠요"}, widths=[12, 80])
body(ws, 2, fixed_cols=(1,), rows=14)
for i, (k, v) in enumerate([
    ("1차 시간", "18:30"), ("1차 장소", "이서식당"), ("1차 주소", "경기 안산시 상록구 한양대학로 68 한대빌딩 2층"), ("1차 지도링크", ""),
    ("2차 시간", "1차 끝나고"), ("2차 장소", "에이쁠"), ("2차 주소", ""), ("2차 지도링크", ""),
    ("회비", ""), ("이동방법", "풋살장에서 이서식당까지 걸어서 약 19분 (1.3km)"), ("공지", ""),
], start=2):
    put(ws, i, [k, v], fixed_cols=(1,))

# ───────────────────────── 설정 ─────────────────────────
ws = wb.create_sheet("설정")
ws.sheet_properties.tabColor = "7F7F7F"
cols = ["항목", "값", "설명"]
header(ws, cols, widths=[18, 22, 60])
body(ws, 3, fixed_cols=(1, 3), rows=10)
for i, (k, v, d) in enumerate([
    ("팀공개", "X", "O로 바꾸면 사이트에 팀과 선수 명단이 공개돼요"),
    ("개회식", "13:00", "개회식 & 팀 제비뽑기 시간 (진행 순서에 표시)"),
    ("시상식", "17:00", "시상식 시간 (진행 순서에 표시)"),
    ("조별리그 경기시간", "12분", "진행 방식 안내에 표시"),
    ("토너먼트 경기시간", "12분 (결승 15분)", "진행 방식 안내에 표시"),
    ("토너먼트 무승부", "승부차기 3명", "진행 방식 안내에 표시"),
], start=2):
    put(ws, i, [k, v, d], fixed_cols=(1, 3))
dropdown(ws, "B", ["O", "X"], ref="B2")  # 팀공개 칸만

wb.active = 0
wb.save(OUT)
print("saved", OUT)
